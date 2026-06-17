"""
Reduce the current V10 deck to the requested 5-slide structure.

The deck is reduced in-place from AI-Driven Embedded Product Lifecycle
OrchestratorV10.pptx:
- keep original page 1 as-is
- replace page 2 with one combined summary slide for pages 2, 3, 4, 5 and 7
- keep original page 6 (JTAG) as-is
- keep original page 8 (STQC) as-is
- keep original page 9 (Results) as-is
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


WORKSPACE = Path("/workspace")
OUTPUT = WORKSPACE / "AI-Driven Embedded Product Lifecycle OrchestratorV10.pptx"

NAVY = RGBColor(0x05, 0x14, 0x2E)
DARK_NAVY = RGBColor(0x02, 0x0A, 0x1A)
BLUE = RGBColor(0x00, 0x72, 0xC6)
CYAN = RGBColor(0x00, 0xC8, 0xF0)
TEAL = RGBColor(0x00, 0xA8, 0x9A)
MID = RGBColor(0x0E, 0x2A, 0x4A)
PANEL = RGBColor(0x0B, 0x1F, 0x3A)
CARD = RGBColor(0x0F, 0x28, 0x45)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xB8, 0xD0, 0xEC)
DIM = RGBColor(0x60, 0x80, 0xA8)
GREEN = RGBColor(0x00, 0xD4, 0x8A)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)
RED = RGBColor(0xFF, 0x3A, 0x3A)
PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
GOLD = RGBColor(0xF5, 0xC5, 0x18)


def delete_slides_except(prs, keep_indices):
    for idx in reversed(range(len(prs.slides))):
        if idx in keep_indices:
            continue
        slide_id = prs.slides._sldIdLst[idx]
        prs.part.drop_rel(slide_id.rId)
        del prs.slides._sldIdLst[idx]


def clear_slide(slide):
    for shape in list(slide.shapes):
        element = shape._element
        element.getparent().remove(element)


def rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def roundrect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(5, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def oval(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(9, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def txt(slide, text, x, y, w, h, size=12, bold=False, italic=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    box.word_wrap = True
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Segoe UI"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def draw_combined_slide(slide):
    clear_slide(slide)
    rect(slide, 0, 0, Inches(10.0), Inches(7.5), DARK_NAVY)
    rect(slide, 0, 0, Inches(10.0), Inches(0.055), CYAN)

    roundrect(slide, Inches(0.38), Inches(0.48), Inches(2.2), Inches(0.3), TEAL)
    txt(slide, "LIFECYCLE FLOW", Inches(0.38), Inches(0.53), Inches(2.2), Inches(0.12), size=8.5, bold=True, align=PP_ALIGN.CENTER)
    txt(slide, "Combined Lifecycle Summary", Inches(0.38), Inches(0.82), Inches(9.0), Inches(0.58), size=28, bold=True)
    rect(slide, Inches(0.38), Inches(1.45), Inches(9.24), Inches(0.016), TEAL)
    txt(
        slide,
        "Pages 2, 3, 4, 5 and 7 compressed into one V9-style flow: problem and architecture, PRD-to-work selection, PG3 generation and QA release gating.",
        Inches(0.38), Inches(1.55), Inches(9.3), Inches(0.45), size=12, italic=True, color=LIGHT,
    )

    steps = [
        (
            TEAL,
            "01",
            "UNIFY",
            "Problem + Architecture",
            "Fragmented PRD, Jira, Confluence, GitHub, QA and hardware evidence are connected into AEPLO's knowledge graph and orchestration core.",
        ),
        (
            BLUE,
            "02",
            "PICK",
            "PRD -> Epics -> Branch",
            "Confluence PRD guided by SDE Elements and PSJIRA creates epics/stories for users to pick; selected work creates a feature branch.",
        ),
        (
            GREEN,
            "03",
            "VALIDATE",
            "PG3 Firmware -> QA Gate",
            "Datasheet + PRD + SDK/BSP + OS/RTOS generate firmware; QA inputs, PRD and JTAG evidence produce READY / NOT_READY.",
        ),
    ]
    cy_flow = Inches(3.62)
    rect(slide, Inches(0.38), cy_flow - Inches(0.015), Inches(9.24), Inches(0.03), RGBColor(0x08, 0x28, 0x48))

    for i, (color, num, verb, subtitle, detail) in enumerate(steps):
        x = Inches(0.38 + i * 3.12)
        y = Inches(2.2)
        card_w = Inches(2.88)
        card_h = Inches(4.52)
        roundrect(slide, x, y, card_w, card_h, CARD)
        rect(slide, x, y, card_w, Inches(0.05), color)
        oval(slide, x + card_w / 2 - Inches(0.42), cy_flow - Inches(0.42), Inches(0.84), Inches(0.84), color)
        txt(slide, num, x + card_w / 2 - Inches(0.42), cy_flow - Inches(0.24), Inches(0.84), Inches(0.16), size=14, bold=True, align=PP_ALIGN.CENTER)
        if i < 2:
            for j in range(6):
                rect(slide, x + card_w + Inches(0.04) + j * Inches(0.028), cy_flow - Inches(0.016), Inches(0.02), Inches(0.032), TEAL)
        txt(slide, verb, x, y + Inches(0.15), card_w, Inches(0.45), size=20, bold=True, color=color, align=PP_ALIGN.CENTER)
        txt(slide, subtitle, x, y + Inches(0.62), card_w, Inches(0.4), size=12, bold=True, align=PP_ALIGN.CENTER)
        rect(slide, x + Inches(0.18), y + Inches(1.08), card_w - Inches(0.36), Inches(0.016), color)
        txt(slide, detail, x + Inches(0.18), y + Inches(1.22), card_w - Inches(0.36), Inches(3.1), size=11.2, color=LIGHT)

    roundrect(slide, Inches(0.35), Inches(7.0), Inches(9.3), Inches(0.34), MID)
    rect(slide, Inches(0.35), Inches(7.0), Inches(0.05), Inches(0.34), TEAL)
    txt(
        slide,
        "Output: users pick traceable work, create feature branches, validate PG3 firmware with JTAG evidence, and loop through QA until READY.",
        Inches(0.55), Inches(7.0), Inches(9.0), Inches(0.34), size=10, bold=True, color=TEAL, align=PP_ALIGN.CENTER,
    )
    txt(slide, "2  /  5", Inches(9.05), Inches(7.17), Inches(0.45), Inches(0.16), size=8.5, color=DIM, align=PP_ALIGN.RIGHT)
    txt(slide, "HONEYWELL", Inches(8.1), Inches(7.17), Inches(0.85), Inches(0.16), size=8.5, bold=True, color=DIM, align=PP_ALIGN.RIGHT)


def main():
    prs = Presentation(OUTPUT)
    if len(prs.slides) >= 9:
        # Keep original pages 1, 2 (converted summary), 6, 8 and 9.
        keep = {0, 1, 5, 7, 8}
        draw_combined_slide(prs.slides[1])
        delete_slides_except(prs, keep)
    elif len(prs.slides) == 5:
        draw_combined_slide(prs.slides[1])
    else:
        raise RuntimeError(f"Expected 9-slide source or 5-slide reduced deck, found {len(prs.slides)} slides")

    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
