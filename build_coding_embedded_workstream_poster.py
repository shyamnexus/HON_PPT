"""
One-page poster candidate for the Coding Embedded Workstream common idea.
Generates: Coding Embedded Workstream - Common Idea Poster.pptx
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


# Slide size: widescreen 16:9
W = Inches(13.333)
H = Inches(7.5)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


# Palette aligned with the AEPLO deck.
DARK = RGBColor(0x02, 0x0A, 0x1A)
NAVY = RGBColor(0x05, 0x14, 0x2E)
PANEL = RGBColor(0x0B, 0x1F, 0x3A)
CARD = RGBColor(0x0F, 0x28, 0x45)
BLUE = RGBColor(0x00, 0x72, 0xC6)
CYAN = RGBColor(0x00, 0xC8, 0xF0)
TEAL = RGBColor(0x00, 0xA8, 0x9A)
GREEN = RGBColor(0x00, 0xD4, 0x8A)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)
PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
RED = RGBColor(0xFF, 0x3A, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xB8, 0xD0, 0xEC)
DIM = RGBColor(0x60, 0x80, 0xA8)


def rect(slide, l, t, w, h, fill, line_color=None):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def roundrect(slide, l, t, w, h, fill, line_color=None):
    shape = slide.shapes.add_shape(5, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def oval(slide, l, t, w, h, fill, line_color=None):
    shape = slide.shapes.add_shape(9, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def txt(slide, text, l, t, w, h, size=14, bold=False, italic=False,
        color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    box.word_wrap = True
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Segoe UI"
    return box


def bullet_list(slide, items, l, t, w, h, size=10.5, color=LIGHT):
    box = slide.shapes.add_textbox(l, t, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    for idx, item in enumerate(items):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.level = 0
        run = paragraph.add_run()
        run.text = "- " + item
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Segoe UI"
    return box


def arrow(slide, x, y, w, color=CYAN):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Inches(0.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def stat_card(slide, l, t, w, value, label, color):
    roundrect(slide, l, t, w, Inches(0.82), PANEL)
    rect(slide, l, t, Inches(0.06), Inches(0.82), color)
    txt(slide, value, l + Inches(0.14), t + Inches(0.08), w - Inches(0.28), Inches(0.28),
        size=19, bold=True, color=color, align=PP_ALIGN.CENTER)
    txt(slide, label, l + Inches(0.14), t + Inches(0.42), w - Inches(0.28), Inches(0.28),
        size=8.8, color=LIGHT, align=PP_ALIGN.CENTER)


def build_poster():
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, W, H, DARK)
    rect(slide, 0, 0, W, Inches(0.08), CYAN)
    rect(slide, 0, H - Inches(0.08), W, Inches(0.08), PURPLE)

    # Decorative embedded/circuit motif.
    for i in range(18):
        rect(slide, Inches(8.8), Inches(0.32 + i * 0.36),
             Inches(4.25), Inches(0.008), RGBColor(0x06, 0x1C, 0x38))
    for i in range(11):
        rect(slide, Inches(8.85 + i * 0.42), Inches(0.32),
             Inches(0.008), Inches(6.45), RGBColor(0x06, 0x1C, 0x38))
    oval(slide, Inches(10.85), Inches(0.78), Inches(1.9), Inches(1.9),
         RGBColor(0x06, 0x1C, 0x38))
    oval(slide, Inches(11.15), Inches(1.08), Inches(1.3), Inches(1.3), PANEL)
    oval(slide, Inches(11.48), Inches(1.41), Inches(0.64), Inches(0.64), NAVY)

    # Header.
    txt(slide, "CODING EMBEDDED WORKSTREAM", Inches(0.35), Inches(0.28),
        Inches(5.7), Inches(0.35), size=13, bold=True, color=CYAN)
    txt(slide, "AEPLO JTAG Accelerator for Embedded PDLC AI Optimisation",
        Inches(0.35), Inches(0.68), Inches(8.95), Inches(0.58),
        size=27, bold=True, color=WHITE)
    txt(slide,
        "Common idea: make JTAG/SWD hardware evidence a native AEPLO capability, "
        "optimising the embedded product development lifecycle from discovery to launch.",
        Inches(0.38), Inches(1.27), Inches(8.7), Inches(0.48),
        size=12.3, italic=True, color=LIGHT)

    roundrect(slide, Inches(9.55), Inches(0.54), Inches(3.35), Inches(0.82), BLUE)
    txt(slide, "NPI GATE DEMO PLAN", Inches(9.55), Inches(0.62),
        Inches(3.35), Inches(0.25), size=12, bold=True,
        color=WHITE, align=PP_ALIGN.CENTER)
    txt(slide, "PG1 Discovery  ->  PG3 Development  ->  PG5 Validation & Launch Prep",
        Inches(9.55), Inches(0.94), Inches(3.35), Inches(0.25), size=8.6, color=CYAN,
        align=PP_ALIGN.CENTER)

    # Three emphasis pillars requested by the workstream.
    pillars = [
        (BLUE, "EMBEDDED", "MCU / RTOS / JTAG-SWD / peripherals / board evidence"),
        (PURPLE, "PDLC", "Requirements -> code -> debug -> validate -> launch"),
        (GREEN, "AI OPTIMISATION", "Risk scoring, RCA, test focus, controlled fixes"),
    ]
    for i, (color, title, body) in enumerate(pillars):
        x = Inches(0.35 + i * 4.1)
        roundrect(slide, x, Inches(1.92), Inches(3.85), Inches(0.82), PANEL)
        rect(slide, x, Inches(1.92), Inches(0.06), Inches(0.82), color)
        txt(slide, title, x + Inches(0.18), Inches(2.02),
            Inches(1.55), Inches(0.25), size=11.5, bold=True, color=color)
        txt(slide, body, x + Inches(1.55), Inches(1.98),
            Inches(2.12), Inches(0.42), size=8.8, color=LIGHT)

    # AEPLO capability statement.
    roundrect(slide, Inches(0.35), Inches(2.95), Inches(12.6), Inches(0.62), CARD)
    rect(slide, Inches(0.35), Inches(2.95), Inches(0.06), Inches(0.62), CYAN)
    txt(slide, "AEPLO CAPABILITY", Inches(0.55), Inches(3.08),
        Inches(1.9), Inches(0.25), size=11.5, bold=True, color=CYAN)
    txt(slide,
        "JTAG Accelerator is part of AEPLO: a multi-agent embedded coding accelerator that links Jira, IDE, Git, "
        "J-Link/JTAG, target MCU evidence, AI root-cause analysis, controlled PRs and release gates.",
        Inches(2.32), Inches(3.06), Inches(10.35), Inches(0.28), size=9.7, color=LIGHT)

    # NPI gate demo flow.
    txt(slide, "DEMO BY NPI GATES", Inches(0.35), Inches(3.78),
        Inches(2.8), Inches(0.28), size=12.5, bold=True, color=WHITE)
    gates = [
        (
            BLUE,
            "PG1",
            "DISCOVERY",
            [
                "Ingest PRD, Jira themes, defect history and target board context.",
                "AI maps PDLC risks, trace gaps and embedded debug hotspots.",
                "Demo output: risk-backed scope, board setup and validation plan.",
            ],
        ),
        (
            ORANGE,
            "PG3",
            "DEVELOPMENT",
            [
                "JTAG Accelerator attaches to IDE, J-Link and target MCU.",
                "Fault Analyzer, Variable Tracer and Peripheral Inspector diagnose coding defects.",
                "Demo output: RCA, guarded patch, PR and register/memory evidence.",
            ],
        ),
        (
            GREEN,
            "PG5",
            "VALIDATION & LAUNCH PREP",
            [
                "AEPLO prioritises tests, checks fix evidence and updates traceability.",
                "Release Gate evaluates code, QA, hardware proof and launch readiness.",
                "Demo output: READY / NOT_READY verdict with action list.",
            ],
        ),
    ]
    for i, (color, gate, title, items) in enumerate(gates):
        x = Inches(0.35 + i * 4.27)
        roundrect(slide, x, Inches(4.12), Inches(3.9), Inches(1.72), CARD)
        rect(slide, x, Inches(4.12), Inches(3.9), Inches(0.06), color)
        oval(slide, x + Inches(0.18), Inches(4.34), Inches(0.6), Inches(0.6), color)
        txt(slide, gate, x + Inches(0.18), Inches(4.49),
            Inches(0.6), Inches(0.18), size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        txt(slide, title, x + Inches(0.92), Inches(4.31),
            Inches(2.75), Inches(0.28), size=12.5, bold=True, color=color)
        bullet_list(slide, items, x + Inches(0.22), Inches(4.82),
                    Inches(3.45), Inches(0.78), size=7.8, color=LIGHT)
        if i < len(gates) - 1:
            arrow(slide, x + Inches(3.95), Inches(4.86), Inches(0.58), color=CYAN)

    # Bottom demo storyline and outcomes.
    roundrect(slide, Inches(0.35), Inches(6.12), Inches(6.45), Inches(0.9), PANEL)
    rect(slide, Inches(0.35), Inches(6.12), Inches(0.06), Inches(0.9), PURPLE)
    txt(slide, "POSTER MESSAGE", Inches(0.55), Inches(6.22),
        Inches(1.75), Inches(0.24), size=11.2, bold=True, color=PURPLE)
    txt(slide,
        "AEPLO optimises embedded PDLC decisions at each NPI gate: discover risk early, "
        "accelerate coding/debug with JTAG evidence, and validate launch readiness with AI gates.",
        Inches(0.55), Inches(6.5), Inches(5.95), Inches(0.28), size=9.2, color=LIGHT)

    txt(slide, "SUCCESS SIGNALS", Inches(7.12), Inches(6.06),
        Inches(2.0), Inches(0.24), size=11.2, bold=True, color=GREEN)
    stat_card(slide, Inches(7.12), Inches(6.36), Inches(1.38), "PG1", "risk clarity", BLUE)
    stat_card(slide, Inches(8.66), Inches(6.36), Inches(1.38), "<30m", "RCA target", ORANGE)
    stat_card(slide, Inches(10.2), Inches(6.36), Inches(1.38), "100%", "evidence link", PURPLE)
    stat_card(slide, Inches(11.74), Inches(6.36), Inches(1.2), "PG5", "launch gate", GREEN)

    # Footer.
    txt(slide, "Common Idea | AEPLO + JTAG Accelerator | Embedded PDLC AI Optimisation",
        Inches(0.35), Inches(7.18), Inches(8.4), Inches(0.18),
        size=8.5, color=DIM)
    txt(slide, "One-page poster candidate for Coding Embedded Workstream",
        Inches(9.15), Inches(7.18), Inches(3.8), Inches(0.18),
        size=8.5, color=DIM, align=PP_ALIGN.RIGHT)


build_poster()

OUTPUT = "/workspace/Coding Embedded Workstream - Common Idea Poster.pptx"
prs.save(OUTPUT)
print(f"Saved: {OUTPUT} ({len(prs.slides)} slide)")
